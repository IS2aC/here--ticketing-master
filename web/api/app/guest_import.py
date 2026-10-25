"""Lecture et vérification d'une liste d'invités (Excel .xlsx ou CSV).

Colonnes attendues (dans n'importe quel ordre, en-têtes reconnus sans tenir compte des accents ni de la casse) :
nom, prénoms, numéro (facultatif), catégorie (facultatif, « Normal » par défaut), position (facultatif : table ou place).
"""

import csv
import enum
import io
import re
import unicodedata
from dataclasses import asdict, dataclass, field
from datetime import datetime
from typing import Any

from openpyxl import load_workbook

from app.phone import normalize_phone

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_ROWS = 5000
DEFAULT_CATEGORY = "Normal"


class ImportFileError(ValueError):
    """Fichier illisible : le message est affiché tel quel à l'organisateur."""


def normalize(text: str) -> str:
    """« Prénom(s) » → « prenom s », « CATÉGORIE » → « categorie »."""
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", " ", text.lower()).strip()


# En-tête normalisé → champ Here!.
_HEADERS = {
    "last_name": {
        "nom",
        "nom de famille",
        "noms",
        "last name",
        "nom complet",
        "invite",
        "nom de l invite",
    },
    "first_names": {"prenom", "prenoms", "prenom s", "first name", "first names"},
    "phone": {
        "numero",
        "numeros",
        "num",
        "telephone",
        "tel",
        "whatsapp",
        "numero whatsapp",
        "numero de telephone",
        "contact",
        "mobile",
        "portable",
        "phone",
    },
    "category": {"categorie", "categories", "groupe", "category", "type"},
    "seat": {
        "position",
        "table",
        "place",
        "placement",
        "numero de table",
        "table place",
    },
}
_HEADER_LOOKUP = {
    alias: field_name for field_name, aliases in _HEADERS.items() for alias in aliases
}

FIELD_LABELS = {
    "last_name": "Nom",
    "first_names": "Prénoms",
    "phone": "Numéro",
    "category": "Catégorie",
    "seat": "Position",
}


@dataclass
class ColumnMatch:
    header: str
    field: str | None  # None = colonne ignorée


@dataclass
class RawRow:
    line: int
    last_name: str | None = None
    first_names: str | None = None
    phone: str | None = None
    category: str | None = None
    seat: str | None = None


@dataclass
class ParsedFile:
    columns: list[ColumnMatch]
    rows: list[RawRow] = field(default_factory=list)


def _cell(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, int):
        digits = str(value)
        # Excel enlève le 0 initial des numéros saisis comme nombres : 0707123456 → 707123456.
        if len(digits) == 9:
            return "0" + digits
        if len(digits) >= 11:
            return "+" + digits
        return digits
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y")
    text = str(value).strip()
    return text or None


def _read_table(file_name: str, content: bytes) -> list[list[Any]]:
    extension = file_name.rsplit(".", 1)[-1].lower() if "." in file_name else ""
    if extension == "xlsx":
        try:
            workbook = load_workbook(
                io.BytesIO(content), read_only=True, data_only=True
            )
        except Exception as exc:
            raise ImportFileError(
                "Nous n'arrivons pas à lire ce fichier Excel. Vérifiez qu'il s'ouvre bien."
            ) from exc
        sheet = workbook.worksheets[0]
        table = [
            list(row)
            for row in sheet.iter_rows(values_only=True, max_row=MAX_ROWS + 20)
        ]
        workbook.close()
        return table
    if extension in {"csv", "txt"}:
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = content.decode(
                "cp1252", errors="replace"
            )  # CSV enregistré par Excel sous Windows
        delimiter = ";" if text.count(";") >= text.count(",") else ","
        try:
            delimiter = csv.Sniffer().sniff(text[:4096], delimiters=";,\t").delimiter
        except csv.Error:
            pass
        return list(csv.reader(io.StringIO(text), delimiter=delimiter))
    if extension == "xls":
        raise ImportFileError(
            "Ce fichier est au format Excel ancien (.xls). Enregistrez-le en .xlsx ou en .csv."
        )
    raise ImportFileError("Choisissez un fichier Excel (.xlsx) ou CSV.")


def parse_file(file_name: str, content: bytes) -> ParsedFile:
    if len(content) > MAX_FILE_BYTES:
        raise ImportFileError(
            "Ce fichier dépasse 5 Mo. Gardez seulement la liste des invités."
        )
    table = _read_table(file_name, content)

    # L'en-tête est la première ligne (parmi les 10 premières) où l'on reconnaît le nom ou les prénoms.
    header_index = None
    for i, row in enumerate(table[:10]):
        fields = {_HEADER_LOOKUP.get(normalize(str(c))) for c in row if c is not None}
        if fields & {"last_name", "first_names"}:
            header_index = i
            break
    if header_index is None:
        raise ImportFileError(
            "Nous n'avons pas reconnu les colonnes. La première ligne doit contenir au moins « Nom » ou « Prénoms »."
        )

    columns: list[ColumnMatch] = []
    used: set[str] = set()
    for cell in table[header_index]:
        header = _cell(cell) or ""
        field_name = _HEADER_LOOKUP.get(normalize(header))
        if field_name in used:
            field_name = None  # deux colonnes « Téléphone » : on garde la première
        if field_name:
            used.add(field_name)
        columns.append(ColumnMatch(header=header, field=field_name))

    parsed = ParsedFile(columns=columns)
    for offset, row in enumerate(table[header_index + 1 :], start=header_index + 2):
        values: dict[str, str | None] = {}
        for column, cell in zip(columns, row, strict=False):
            if column.field:
                values[column.field] = _cell(cell)
        if not any(values.values()):
            continue
        if len(parsed.rows) >= MAX_ROWS:
            raise ImportFileError(
                f"Ce fichier contient plus de {MAX_ROWS} invités. Découpez-le en plusieurs fichiers."
            )
        parsed.rows.append(RawRow(line=offset, **values))
    if not parsed.rows:
        raise ImportFileError(
            "Ce fichier ne contient aucun invité sous la ligne des titres."
        )
    return parsed


# --- Vérification des lignes -------------------------------------------------------------


class RowStatus(enum.StrEnum):
    ready = "ready"  # importée
    phone_missing = "phone_missing"  # importée, invitation impossible sans numéro
    phone_invalid = "phone_invalid"  # importée, numéro à corriger
    name_missing = "name_missing"  # non importée
    category_unknown = "category_unknown"  # non importée
    duplicate = "duplicate"  # non importée (même numéro déjà présent)


IMPORTED = {RowStatus.ready, RowStatus.phone_missing, RowStatus.phone_invalid}


@dataclass
class CheckedRow:
    line: int
    display_name: str
    last_name: str | None
    first_names: str | None
    phone: str | None  # normalisé si valide, sinon tel que saisi
    phone_valid: bool
    category: str | None  # nom de la catégorie de l'événement
    seat: str | None
    status: RowStatus
    raw_category: str | None = None


def guest_key(phone: str | None, display_name: str) -> str:
    """Identité d'un invité pour repérer les doublons : son numéro, ou à défaut son nom."""
    return phone if phone else "nom:" + normalize(display_name)


def check_rows(
    rows: list[RawRow], categories: list[str], existing_keys: set[str]
) -> list[CheckedRow]:
    """Vérifie chaque ligne.

    `categories` : catégories de l'événement ; `existing_keys` : invités déjà dans la liste (voir `guest_key`).
    """
    by_key = {normalize(name): name for name in categories}
    seen = set(existing_keys)
    checked: list[CheckedRow] = []
    for row in rows:
        display_name = " ".join(p for p in (row.first_names, row.last_name) if p)
        phone = normalize_phone(row.phone) if row.phone else None
        stored_phone = phone or row.phone

        raw_category = row.category
        if raw_category:
            category = by_key.get(normalize(raw_category))
        else:
            category = by_key.get(normalize(DEFAULT_CATEGORY), DEFAULT_CATEGORY)

        if not display_name:
            status = RowStatus.name_missing
        elif category is None:
            status = RowStatus.category_unknown
        elif guest_key(stored_phone, display_name) in seen:
            status = RowStatus.duplicate
        elif not row.phone:
            status = RowStatus.phone_missing
        elif phone is None:
            status = RowStatus.phone_invalid
        else:
            status = RowStatus.ready
        if status in IMPORTED:
            seen.add(guest_key(stored_phone, display_name))

        checked.append(
            CheckedRow(
                line=row.line,
                display_name=display_name,
                last_name=row.last_name,
                first_names=row.first_names,
                phone=stored_phone,
                phone_valid=phone is not None,
                category=category,
                seat=row.seat,
                status=status,
                raw_category=raw_category,
            )
        )
    return checked


def rows_to_json(rows: list[RawRow]) -> list[dict[str, Any]]:
    return [asdict(r) for r in rows]


def rows_from_json(data: list[dict[str, Any]]) -> list[RawRow]:
    return [RawRow(**r) for r in data]
